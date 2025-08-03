# tools/tool_secretary.py
import logging
import tempfile
from langchain.agents import Tool
from docx import Document as WordDocument
from openpyxl import Workbook as ExcelWorkbook
from fpdf import FPDF

logger = logging.getLogger(__name__)

def create_document(input_str: str) -> str:
    logger.info(f"Эксперт 'Secretary': Получена задача на создание документа.")
    try:
        parts = input_str.split('|')
        if len(parts) != 2:
            return "Ошибка: неверный формат. Используйте 'текст|тип документа' (word, excel, pdf)."
        content, doc_type = parts[0].strip(), parts[1].strip().lower()

        if doc_type == 'word':
            doc = WordDocument()
            doc.add_paragraph(content)
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
            doc.save(temp_file.name)
            return f"Документ Word успешно создан: {temp_file.name}"
        elif doc_type == 'excel':
            wb = ExcelWorkbook()
            ws = wb.active
            for line in content.split('\n'):
                ws.append(line.split(','))
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
            wb.save(temp_file.name)
            return f"Документ Excel успешно создан: {temp_file.name}"
        elif doc_type == 'pdf':
            pdf = FPDF()
            pdf.add_page()
            pdf.add_font('DejaVu', '', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', uni=True)
            pdf.set_font('DejaVu', '', 12)
            pdf.multi_cell(0, 10, content)
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
            pdf.output(temp_file.name)
            return f"PDF документ успешно создан: {temp_file.name}"
        else:
            return "Неподдерживаемый тип документа."
    except Exception as e:
        return f"Ошибка при создании документа: {e}"

secretary_tool = Tool(
    name="Secretary",
    func=create_document,
    description="Используй для создания документов. Входные данные: 'текст|тип документа'."
)